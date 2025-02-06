import os
import pytest
from docx import Document
from extract import TextExtractor, DocxTextExtractor, PdfTextExtractor, get_text_extractor, extract_text

# Helper function to create a temporary DOCX file
def create_temp_docx(file_path, content):
    doc = Document()
    doc.add_paragraph(content)
    doc.save(file_path)

# Helper function to create a temporary PDF file
def create_temp_pdf(file_path, content):
    import pdfplumber
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt=content, ln=True, align="C")
    pdf.output(file_path)

# Test case for valid DOCX file
def test_docx_text_extractor(tmpdir):
    file_path = os.path.join(tmpdir, "test.docx")
    content = "This is a test DOCX file."
    create_temp_docx(file_path, content)

    extractor = DocxTextExtractor()
    extracted_text = extractor.extract_text(file_path)
    assert extracted_text == content.lower()

# Test case for valid PDF file
def test_pdf_text_extractor(tmpdir):
    file_path = os.path.join(tmpdir, "test.pdf")
    content = "This is a test PDF file."
    create_temp_pdf(file_path, content)

    extractor = PdfTextExtractor()
    extracted_text = extractor.extract_text(file_path)
    assert extracted_text == content.lower()

# Test case for unsupported file format
def test_unsupported_file_format(tmpdir):
    file_path = os.path.join(tmpdir, "test.txt")
    with open(file_path, "w") as f:
        f.write("This is a test TXT file.")

    extractor = get_text_extractor(file_path)
    assert extractor is None

# Test case for invalid file path
def test_invalid_file_path():
    file_path = "nonexistent.docx"
    extractor = DocxTextExtractor()
    extracted_text = extractor.extract_text(file_path)
    assert extracted_text == ""

# Test case for empty DOCX file
def test_empty_docx_file(tmpdir):
    file_path = os.path.join(tmpdir, "empty.docx")
    create_temp_docx(file_path, "")

    extractor = DocxTextExtractor()
    extracted_text = extractor.extract_text(file_path)
    assert extracted_text == ""

# Test case for empty PDF file
def test_empty_pdf_file(tmpdir):
    file_path = os.path.join(tmpdir, "empty.pdf")
    create_temp_pdf(file_path, "")

    extractor = PdfTextExtractor()
    extracted_text = extractor.extract_text(file_path)
    assert extracted_text == ""

# Test case for factory function with DOCX file
def test_get_text_extractor_docx():
    file_path = "test.docx"
    extractor = get_text_extractor(file_path)
    assert isinstance(extractor, DocxTextExtractor)

# Test case for factory function with PDF file
def test_get_text_extractor_pdf():
    file_path = "test.pdf"
    extractor = get_text_extractor(file_path)
    assert isinstance(extractor, PdfTextExtractor)

# Test case for main extract_text function with DOCX file
def test_extract_text_docx(tmpdir):
    file_path = os.path.join(tmpdir, "test.docx")
    content = "This is a test DOCX file."
    create_temp_docx(file_path, content)

    extracted_text = extract_text(file_path)
    assert extracted_text == content.lower()

# Test case for main extract_text function with PDF file
def test_extract_text_pdf(tmpdir):
    file_path = os.path.join(tmpdir, "test.pdf")
    content = "This is a test PDF file."
    create_temp_pdf(file_path, content)

    extracted_text = extract_text(file_path)
    assert extracted_text == content.lower()